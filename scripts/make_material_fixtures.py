#!/usr/bin/env python3
"""生成纯合成材料，绝不覆盖既有文件。

python make_material_fixtures.py /path/to/new-or-empty-isolated-directory

依赖仅Pillow、pypdf、python-docx、python-pptx及标准库。
XLSX用固定OOXML制作，以保留“无单元格/显式空串/空格/缺公式缓存”等差异，
不启动Office、不计算公式。不会访问网络、真实业务数据或系统字体文件。
数字只写入图片像素；visual_answers.json是独立答案表，视觉验收前不要给模型。
对外JSON只打印文件清单，不打印答案。build_fixture_bytes()供单元测试复用。
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import secrets
import sys
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile


def repack(parts: dict[str, bytes | str]) -> bytes:
    out = io.BytesIO()
    with ZipFile(out, "w", compression=ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content.encode("utf-8") if isinstance(content, str) else content)
    return out.getvalue()


def replace_zip_part(raw: bytes, part: str, content: bytes | str) -> bytes:
    with ZipFile(io.BytesIO(raw)) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts[part] = content
    return repack(parts)


def _digits_image(number: str) -> bytes:
    from PIL import Image, ImageDraw
    image = Image.new("RGB", (1800, 600), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, 1770, 570), outline="black", width=5)
    # 矢量七段数码，无字体/元数据依赖。
    segments = {"0":"abcedf", "1":"bc", "2":"abged", "3":"abgcd", "4":"fgbc",
                "5":"afgcd", "6":"afgecd", "7":"abc", "8":"abcdefg", "9":"abfgcd"}
    boxes = {"a":(25,0,145,20), "b":(145,20,165,180), "c":(145,200,165,360),
             "d":(25,360,145,380), "e":(5,200,25,360), "f":(5,20,25,180), "g":(25,180,145,200)}
    for i, digit in enumerate(number):
        x, y = 190+i*240, 110
        for segment in segments[digit]:
            x1,y1,x2,y2 = boxes[segment]
            draw.rectangle((x+x1,y+y1,x+x2,y+y2), fill="black")
    buf = io.BytesIO()
    image.save(buf, "PNG")
    image.close()
    return buf.getvalue()


def _pdf(png: bytes) -> bytes:
    from PIL import Image
    from pypdf import PdfWriter
    from pypdf.generic import (ArrayObject, DictionaryObject, NameObject,
                              NumberObject, DecodedStreamObject, TextStringObject)
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    western = DictionaryObject({NameObject("/Type"):NameObject("/Font"),
                               NameObject("/Subtype"):NameObject("/Type1"), NameObject("/BaseFont"):NameObject("/Helvetica")})
    cidinfo = DictionaryObject({NameObject("/Registry"): TextStringObject("Adobe"),
                               NameObject("/Ordering"): TextStringObject("GB1"), NameObject("/Supplement"): NumberObject(4)})
    descendant = DictionaryObject({NameObject("/Type"):NameObject("/Font"), NameObject("/Subtype"):NameObject("/CIDFontType0"),
                                   NameObject("/BaseFont"):NameObject("/STSong-Light"), NameObject("/CIDSystemInfo"):cidinfo})
    chinese_text = "合成材料 中文文本 日期与表格"
    chars = sorted(set(chinese_text))
    mappings = "\n".join(f"<{ord(c):04X}> <{ord(c):04X}>" for c in chars)
    cmap = DecodedStreamObject()
    cmap.set_data(("/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n"
                   "/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
                   "/CMapName /FixtureUnicode def\n/CMapType 2 def\n1 begincodespacerange\n<0000> <FFFF>\nendcodespacerange\n"
                   f"{len(chars)} beginbfchar\n{mappings}\nendbfchar\nendcmap\nCMapName currentdict /CMap defineresource pop\nend\nend\n").encode("ascii"))
    chinese = DictionaryObject({NameObject("/Type"):NameObject("/Font"), NameObject("/Subtype"):NameObject("/Type0"),
                                 NameObject("/BaseFont"):NameObject("/STSong-Light"), NameObject("/Encoding"):NameObject("/UniGB-UCS2-H"),
                                 NameObject("/DescendantFonts"):ArrayObject([writer._add_object(descendant)]),
                                 NameObject("/ToUnicode"):writer._add_object(cmap)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):DictionaryObject({
        NameObject("/F1"):writer._add_object(western), NameObject("/F2"):writer._add_object(chinese)})})
    content = DecodedStreamObject()
    content.set_data(("BT /F1 20 Tf 55 775 Td (Material fixture - text layer) Tj ET\n"
                      f"BT /F2 18 Tf 55 725 Td <{chinese_text.encode('utf-16-be').hex().upper()}> Tj ET\n"
                      "BT /F1 12 Tf 55 675 Td (Date: 2026-09-08   Value: 42) Tj ET\n"
                      "0.5 w 55 550 480 70 re S 55 585 m 535 585 l S 295 550 m 295 620 l S\n"
                      "BT /F1 12 Tf 65 600 Td (Item) Tj 240 0 Td (Value) Tj ET\n"
                      "BT /F1 12 Tf 65 565 Td (Synthetic) Tj 240 0 Td (42) Tj ET\n").encode("ascii"))
    page[NameObject("/Contents")] = writer._add_object(content)
    scan = writer.add_blank_page(width=595, height=842)
    with Image.open(io.BytesIO(png)) as image:
        image = image.convert("RGB")
        obj = DecodedStreamObject()
        obj.set_data(image.tobytes())
        obj.update({NameObject("/Type"):NameObject("/XObject"),NameObject("/Subtype"):NameObject("/Image"),
                    NameObject("/Width"):NumberObject(image.width), NameObject("/Height"):NumberObject(image.height),
                    NameObject("/ColorSpace"):NameObject("/DeviceRGB"),NameObject("/BitsPerComponent"):NumberObject(8)})
        xobj = obj.flate_encode()
    scan[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"):DictionaryObject({NameObject("/Im0"):writer._add_object(xobj)})})
    contents = DecodedStreamObject()
    contents.set_data(b"q 535 0 0 178.333 30 330 cm /Im0 Do Q\n")
    scan[NameObject("/Contents")] = writer._add_object(contents)
    writer.add_outline_item("正文与中文文本", 0)
    writer.add_outline_item("仅图像页：请目视读取", 1)
    writer.add_metadata({"/Title":"合成解析验收材料", "/Producer":"material-fixtures; synthetic only"})
    out = io.BytesIO()
    writer.write(out)
    writer.close()
    return out.getvalue()


def _xlsx() -> bytes:
    ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkgrel = "http://schemas.openxmlformats.org/package/2006/relationships"
    serial = (date(2026,9,8)-date(1899,12,30)).days
    def inline(addr: str, text: str) -> str:
        return f'<c r="{addr}" t="inlineStr"><is><t xml:space="preserve">{escape(text)}</t></is></c>'
    sheet = f'''<worksheet xmlns="{ns}"><dimension ref="A1:F7"/><sheetData>
<row r="1">{inline("A1","名称")}{inline("B1","原值")}{inline("C1","公式")}{inline("D1","空值")}{inline("E1","日期")}{inline("F1","类型")}</row>
<row r="2">{inline("A2","中文测试")}<c r="B2"><v>3</v></c><c r="C2"><f>B2*2</f><v>6</v></c>{inline("D2","")}<c r="E2" s="1"><v>{serial}</v></c><c r="F2" t="b"><v>1</v></c></row>
<row r="3"><c r="A3" t="s"><v>0</v></c><c r="B3"><v>4</v></c><c r="C3"><f>B3*2</f></c><c r="E3"/><c r="F3" t="b"><v>0</v></c></row>
<row r="4">{inline("A4","空格原样保留")}{inline("D4","  ")}<c r="C4" t="str"><f>TEXT(B2,"0")</f><v>3</v></c><c r="E4" s="1"><v>60</v></c><c r="F4" t="e"><v>#DIV/0!</v></c></row>
<row r="5">{inline("A5","=literal_not_formula")}<c r="C5"><f t="shared" si="0" ref="C5:C6">B2*2</f><v>6</v></c><c r="E5" s="1"><v>0.5</v></c></row>
<row r="6"><c r="C6"><f t="shared" si="0"/><v>8</v></c></row>
<row r="7">{inline("A7","合并格")}</row></sheetData><mergeCells count="1"><mergeCell ref="A7:B7"/></mergeCells></worksheet>'''
    return repack({
        "[Content_Types].xml": f'''<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/><Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/><Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>''',
        "_rels/.rels": f'<Relationships xmlns="{pkgrel}"><Relationship Id="r1" Type="{rel}/officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{ns}" xmlns:r="{rel}"><workbookPr date1904="0"/><sheets><sheet name="Sheet1" sheetId="1" r:id="r1"/><sheet name="隐藏" sheetId="2" state="hidden" r:id="r2"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{pkgrel}"><Relationship Id="r1" Type="{rel}/worksheet" Target="worksheets/sheet1.xml"/><Relationship Id="r2" Type="{rel}/worksheet" Target="worksheets/sheet2.xml"/><Relationship Id="r3" Type="{rel}/styles" Target="styles.xml"/><Relationship Id="r4" Type="{rel}/sharedStrings" Target="sharedStrings.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": sheet,
        "xl/worksheets/sheet2.xml": f'<worksheet xmlns="{ns}"><dimension ref="A1"/><sheetData/></worksheet>',
        "xl/sharedStrings.xml": f'<sst xmlns="{ns}" count="1" uniqueCount="1"><si><t>共享中文</t></si></sst>',
        "xl/styles.xml": f'<styleSheet xmlns="{ns}"><fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts><fills count="2"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill></fills><borders count="1"><border/></borders><cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs><cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/><xf numFmtId="14" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1"/></cellXfs><cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles></styleSheet>',
    })


def _docx(png: bytes) -> bytes:
    from docx import Document
    from docx.shared import Inches
    doc = Document()
    doc.add_heading("合成材料：正文与表格", level=1)
    doc.add_paragraph("段落一：中文、English、空格  保持原样。")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    for row, values in zip(table.rows, [["字段", "数值"], ["样本", "42"]]):
        for cell, value in zip(row.cells, values): cell.text = value
    doc.add_paragraph("段落二：位于表格之后。")
    doc.add_picture(io.BytesIO(png), width=Inches(5.5))
    doc.sections[0].header.paragraphs[0].text = "页眉：不属于正文读取范围"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx(png: bytes) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "中文幻灯片：合成材料"
    box = slide.shapes.add_textbox(Inches(.7), Inches(1.3), Inches(8.5), Inches(.8))
    box.text_frame.text = "正文在XML形状顺序中保留；不推断视觉顺序。"
    table = slide.shapes.add_table(2, 2, Inches(.7), Inches(2.3), Inches(8.3), Inches(1.3)).table
    for ri, row in enumerate([["项目", "值"], ["测试", "42"]]):
        for ci, value in enumerate(row): table.cell(ri, ci).text = value
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "仅图像数字：用图片通道验收"
    slide2.shapes.add_picture(io.BytesIO(png), Inches(.5), Inches(1.5), width=Inches(9))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def build_fixture_bytes() -> tuple[dict[str, bytes], dict]:
    from PIL import Image
    from pypdf import PdfReader, PdfWriter
    number = str(secrets.randbelow(900000)+100000)
    png = _digits_image(number)
    pdf, xlsx, docx, pptx = _pdf(png), _xlsx(), _docx(png), _pptx(png)
    files = {"sample.pdf":pdf, "sample.xlsx":xlsx, "sample.docx":docx,
             "sample.pptx":pptx, "visual.png":png,
             "sample.csv":'名称,值,说明\r\n中文,42,"双行\n内容"\r\n空串,,"=只作文本"\r\n'.encode("utf-8"),
             "sample-gbk.csv":"名称,值\r\n中文,42\r\n".encode("gbk"),
             "sample-utf16.csv":"名称,值\r\n中文,42\r\n".encode("utf-16"),
             "broken.pdf":b"%PDF-1.4\nnot a valid PDF\n",
             "disguised.xlsx":b"not a zip or workbook"}
    with Image.open(io.BytesIO(png)) as img:
        for fmt, suffix in [("JPEG","jpg"),("WEBP","webp")]:
            out = io.BytesIO(); img.save(out, fmt); files[f"visual.{suffix}"] = out.getvalue()
    writer = PdfWriter()
    writer.clone_document_from_reader(PdfReader(io.BytesIO(pdf)))
    writer.encrypt("fixture-password")
    buf = io.BytesIO(); writer.write(buf); writer.close()
    files["encrypted.pdf"] = buf.getvalue()
    with ZipFile(io.BytesIO(xlsx)) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
    files["forged-dimension.xlsx"] = replace_zip_part(xlsx, "xl/worksheets/sheet1.xml", sheet.replace('ref="A1:F7"','ref="A1:XFD1048576"',1))
    # 以下恶意样本完全合成，不含可执行载荷或真实外链。
    files["unsafe-dtd.docx"] = replace_zip_part(docx, "word/document.xml", '<!DOCTYPE doc [<!ENTITY fixture "blocked">]><doc>&fixture;</doc>')
    files["unsafe-macro.docx"] = replace_zip_part(docx, "word/vbaProject.bin", b"SYNTHETIC-NOT-EXECUTABLE")
    answers = {"visualDigits": number, "visualFiles": ["visual.png", "visual.jpg", "visual.webp"],
               "pdfImagePage": 2, "warning": "此答案表与图片输入隔离；数字未写入图片元数据或PDF文本层。"}
    return files, answers


def make_fixtures(directory: str | Path) -> dict:
    target = Path(directory)
    if target.is_symlink():
        raise FileExistsError("拒绝以符号链接作为夹具输出目录。")
    if target.exists() and (not target.is_dir() or any(target.iterdir())):
        raise FileExistsError("输出目录须不存在或为空；不覆盖现有文件。")
    files, answers = build_fixture_bytes()
    target.mkdir(parents=True, exist_ok=True)
    manifest = {"syntheticOnly":True, "files":[], "visualAnswerFile":"visual_answers.json"}
    # 独占创建；并发冲突也不会覆盖。失败可能保留已生成的部分文件，调用者可据此检查。
    for name, data in files.items():
        with (target/name).open("xb") as stream: stream.write(data)
        manifest["files"].append({"name":name, "byteLength":len(data), "sha256":hashlib.sha256(data).hexdigest()})
    for name, payload in [("visual_answers.json",answers),("fixture_manifest.json",manifest)]:
        with (target/name).open("x",encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", help="调用者指定的新目录或空目录")
    args = parser.parse_args()
    try:
        manifest = make_fixtures(args.directory)
    except (OSError, ImportError) as exc:
        print(json.dumps({"ok":False,"error":type(exc).__name__,"message":"未覆盖文件；请检查空目录与依赖。"},ensure_ascii=False),file=sys.stderr)
        return 2
    print(json.dumps({"ok":True,**manifest},ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
